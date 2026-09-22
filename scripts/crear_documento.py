from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(__file__).resolve().parents[1]
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)
s.header_distance = s.footer_distance = Inches(.492)
# Preset: compact_reference_guide; header: memo_masthead (sin regla).
# Overrides: Title 22 pt; Code Consolas 9 pt, interlineado 1.0, después 0.
for name, size, color, before, after in [
    ('Normal',11,'182238',0,6), ('Title',22,'182238',0,10),
    ('Subtitle',11,'667085',0,6), ('Heading 1',16,'2E74B5',18,10),
    ('Heading 2',13,'2E74B5',14,7), ('Heading 3',12,'1F4D78',10,5)]:
    st=doc.styles[name]; st.font.name='Calibri'; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color)
    pf=st.paragraph_format; pf.space_before=Pt(before); pf.space_after=Pt(after); pf.line_spacing=1.25
    if name.startswith('Heading'): pf.keep_with_next=True
st=doc.styles.add_style('Codigo',WD_STYLE_TYPE.PARAGRAPH)
st.font.name='Consolas'; st.font.size=Pt(9)
st.paragraph_format.space_before=Pt(0); st.paragraph_format.space_after=Pt(0); st.paragraph_format.line_spacing=1
st=doc.styles.add_style('Pregunta',WD_STYLE_TYPE.PARAGRAPH)
st.base_style=doc.styles['Normal']; st.font.bold=True; st.paragraph_format.keep_with_next=True
s.header.paragraphs[0].text='UTP  /  DESARROLLO WEB INTEGRADO  /  LABORATORIO 07'
s.header.paragraphs[0].runs[0].font.size=Pt(8)
p=s.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
p.add_run('Acosta Loyola · U22205593  |  ')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); p._p.append(fld)
for r in p.runs: r.font.size=Pt(8)
def p(t,style=None): return doc.add_paragraph(t,style)
def h(t): doc.add_heading(t,1)
def code(t):
    for line in t.splitlines(): p(line,'Codigo')
def page(): doc.add_page_break()

p('UNIVERSIDAD TECNOLÓGICA DEL PERÚ','Subtitle')
p('MVC y DAO en arquitectura Java EE','Title')
p('Ficha de trabajo desarrollada · Laboratorio práctico N.° 7','Subtitle')
h('I. Datos generales')
p('Facultad de Ingeniería · Carrera de Ingeniería de Sistemas e Informática')
p('Asignatura: Desarrollo Web Integrado (1000005154)\nUnidad 2: Aplicaciones Web y Patrones de Diseño\nSesión: Semana 7 – Sesión 13\nDocente: Mtro. Ing. César A. Cárdenas Latorre\nEstudiante: Acosta Loyola, Jorge Eduardo\nCódigo: U22205593\nFecha de la ficha: 21/09/2026')
h('II. Logro de la sesión')
p('Implementar los patrones MVC y DAO para estructurar una aplicación web, separando la interfaz de usuario, la lógica de control y la persistencia de datos.')
h('III. Actividad 1: control conceptual')
p('1. ¿Por qué es una mala práctica escribir SQL y conexiones JDBC dentro de una JSP o un Servlet?','Pregunta')
p('Porque mezcla presentación o control con persistencia, aumenta el acoplamiento y dificulta las pruebas y el mantenimiento. El SQL y la administración de recursos JDBC deben encapsularse en el DAO; el Servlet coordina solicitudes y la JSP presenta datos.')
p('2. Responsabilidad principal del patrón DAO.','Pregunta')
p('Encapsular el acceso a la fuente de datos y ofrecer operaciones de persistencia mediante una interfaz independiente de la presentación y del controlador.')
p('3. Ventaja de Expression Language y JSTL frente a scriptlets.','Pregunta')
p('EL accede a atributos y propiedades del modelo, mientras JSTL expresa iteraciones y condiciones de forma declarativa; así la vista resulta más legible y evita incrustar lógica Java. Para mostrar datos externos se utiliza c:out, que escapa HTML; EL por sí sola no lo hace.')

page(); h('IV. Actividad 2: código e integración')
doc.add_heading('Paso A. Servlet controlador',2)
p('Las cuatro líneas que completan la ficha son:')
code('ClienteDAO dao = new ClienteDAOImpl();\nList<Cliente> lista = dao.listarClientes();\nrequest.setAttribute("listaClientes", lista);\nrequest.getRequestDispatcher("listarClientes.jsp")\n       .forward(request, response);')
p('La variable se declara con la interfaz ClienteDAO y se instancia con ClienteDAOImpl: el controlador depende del contrato de persistencia y permite sustituir su implementación.')
doc.add_heading('Paso B. Vista JSP',2)
p('Respuesta literal a los espacios en blanco de la ficha:')
code('<%@ taglib prefix="c"\n    uri="http://java.sun.com/jsp/jstl/core" %>\n<tbody>\n    <c:forEach var="c" items="${listaClientes}">\n        <tr>\n            <td>${c.id}</td>\n            <td>${c.nombre}</td>\n            <td>${c.email}</td>\n        </tr>\n    </c:forEach>\n</tbody>')
doc.add_heading('Adaptación del programa entregado',2)
p('El equipo tiene Apache Tomcat 11.0.24 y JDK 21. El proyecto usa imports jakarta.servlet.* y la URI jakarta.tags.core. La vista se ubica en /WEB-INF/views/listarClientes.jsp para impedir el acceso directo sin pasar por el controlador. Los valores se muestran con c:out para evitar que contenido ingresado se interprete como HTML.')
p('listarClientes() declara SQLException; el Servlet la captura y responde con un mensaje controlado y estado HTTP 503 ante un problema de conexión. El navegador envía GET para listar y POST para registrar. Después de guardar, el controlador redirige con HTTP 303 para evitar repetir el INSERT al actualizar la página.')

page(); h('V. Actividad 3: registro con DAO')
doc.add_heading('1. Firma en ClienteDAO',2)
code('boolean registrarCliente(Cliente c);')
doc.add_heading('2. Implementación completada de la ficha',2)
code('''@Override
public boolean registrarCliente(Cliente c) {
    boolean estado = false;
    String sql = "INSERT INTO tb_cliente (nombre, email) "
               + "VALUES (?, ?)";
    try (Connection con = ConectaDB.getConexion();
         PreparedStatement ps = con.prepareStatement(sql)) {
        ps.setString(1, c.getNombre());
        ps.setString(2, c.getEmail());
        if (ps.executeUpdate() > 0) {
            estado = true;
        }
    } catch (SQLException e) {
        e.printStackTrace();
    }
    return estado;
}''')
p('Los espacios corresponden a ps.setString(1, c.getNombre()) y ps.setString(2, c.getEmail()). Los índices de parámetros empiezan en 1. PreparedStatement separa los valores de la instrucción SQL y try-with-resources cierra automáticamente la conexión y la sentencia.')
p('En el código ejecutable se sustituye printStackTrace() por un registro del estado SQL en el log del servidor, sin mostrar detalles internos al usuario. La tabla impone UNIQUE(email); un duplicado hace que registrarCliente devuelva false y el controlador informa que no se completó el registro.')
doc.add_heading('Estructura de persistencia',2)
code('''Base de datos: utp_clientes_s7
Tabla: tb_cliente
id      INT AUTO_INCREMENT PRIMARY KEY
nombre  VARCHAR(100) NOT NULL
email   VARCHAR(150) NOT NULL UNIQUE''')
p('La aplicación valida nombre obligatorio, longitud máxima y formato de correo antes de invocar el DAO. Las credenciales se cargan desde un archivo externo en conf de Tomcat y no se distribuyen en el WAR ni en el ZIP.')

page(); h('VI. Ticket de salida')
p('1. ¿Qué parte resultó más desafiante?','Pregunta')
p('Reflexión propuesta para revisión personal: La parte más desafiante fue comprender cómo viaja la lista desde el DAO al Servlet y de este a la JSP mediante un atributo de request. También fue necesario distinguir el contrato de la interfaz de su implementación y adaptar los imports a la versión de Tomcat instalada.')
p('2. Código estructurado y evidencia de funcionamiento.','Pregunta')
p('Se entrega el código separado en paquetes model, dao, config y web, la vista JSP, los estilos, el script SQL y el WAR compilado. Las pruebas se ejecutaron contra Tomcat y MySQL reales. No se realizó la entrega a Canvas.')
doc.add_heading('Pruebas ejecutadas',2)
for t in [
 'Listado: HTTP 200 y contenido JSP/JSTL generado correctamente.',
 'Registro: HTTP 303 y cliente persistido, verificado directamente mediante JDBC.',
 'Caracteres especiales: acentos, comilla y etiquetas guardados como texto; HTML escapado al mostrarlo.',
 'Duplicados: HTTP 409; no se agrega una segunda fila con el mismo correo.',
 'Datos inválidos: HTTP 400; no se insertan filas.',
 'Token de formulario incorrecto: HTTP 403; solicitud rechazada.',
 'Limpieza: se eliminó únicamente el registro temporal generado por la prueba.'
]: p(t)
doc.add_heading('Captura para Canvas: pendiente',2)
p('La conexión de navegador no estuvo disponible durante la ejecución, por lo que no se generó una captura visual ni se verificó visualmente la interfaz. La evidencia funcional reproducible está en evidencias/pruebas.txt. Para completar la evidencia solicitada por la ficha, abrir la dirección siguiente y capturar la tabla con Win + Shift + S:')
code('http://localhost:8080/utp-clientes/listarClientes')

page(); h('Anexo. Ejecución y entrega')
doc.add_heading('En este equipo',2)
p('La aplicación está desplegada en el Tomcat instalado y utiliza MySQL local. Para volver a iniciar Tomcat, ejecutar desde la carpeta del proyecto:')
code('powershell -ExecutionPolicy Bypass -File scripts/iniciar-tomcat.ps1')
doc.add_heading('Instalación desde el programa entregado',2)
p('Descomprimir el ZIP. Con JDK 17 o superior, Tomcat 11 y MySQL disponibles, ejecutar los siguientes comandos en PowerShell. El primer comando solicita la clave de MySQL sin mostrarla:')
code('powershell -ExecutionPolicy Bypass -File scripts/configurar.ps1\npowershell -ExecutionPolicy Bypass -File scripts/desplegar.ps1\npowershell -ExecutionPolicy Bypass -File scripts/verificar.ps1')
p('Los scripts incluyen un parámetro TomcatHome para cambiar la ruta. Las bibliotecas están incluidas para compilar sin Maven. El pom.xml permite importar el proyecto en un IDE o compilar con Maven si está instalado.')
doc.add_heading('Flujo de la aplicación',2)
code('GET /listarClientes\n  -> ClienteServlet -> ClienteDAO -> MySQL\n  -> atributo listaClientes -> JSP -> tabla HTML\n\nPOST /registrarCliente\n  -> validación -> registrarCliente -> INSERT\n  -> redirección 303 -> listado actualizado')
doc.add_heading('Referencias técnicas',2)
p('Apache Tomcat. Guía de migración a Tomcat 11: Java 17 o superior y especificaciones Jakarta.\nhttps://tomcat.apache.org/migration-11.0.html')
p('Eclipse Foundation. Jakarta Tags 3.0: biblioteca Core y URI jakarta.tags.core.\nhttps://jakarta.ee/specifications/tags/3.0/tagdocs/c/tld-summary')
doc.core_properties.title='Laboratorio 7 – MVC y DAO – Acosta Loyola'
doc.core_properties.author='Jorge Eduardo Acosta Loyola'
out=ROOT/'entrega/Ficha_Laboratorio_7_MVC_DAO_Acosta_Loyola.docx'
doc.save(out)
print(out)
